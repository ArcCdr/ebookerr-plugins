"""Tests for DocxDownloadSourcePlugin (DOCX downloads converted to EPUB)."""

from __future__ import annotations

import logging
import re
from pathlib import Path
from typing import Any

import pytest
import requests
import responses
from docx_download_source.plugin import DocxDownloadSourcePlugin
from ebookerr_sdk.domain.ids import make_book_id
from ebookerr_sdk.epub import EpubDocument
from ebookerr_sdk.spi import BookPatch, BookView
from ebookerr_sdk.testing import FakeContext


def _make_docx(
    tmp_path: Path,
    *,
    paras: list[tuple[str | None, str]],
    title: str | None = None,
    author: str | None = None,
) -> Path:
    """Create a test DOCX file with given paragraphs, title, and author.

    Args:
        tmp_path: Temporary directory path.
        paras: List of (style, text) tuples. style=None for normal paragraph.
        title: Optional document title property.
        author: Optional document author property.

    Returns:
        Path to the created DOCX file.
    """
    import docx

    d = docx.Document()
    if title:
        d.core_properties.title = title
    if author:
        d.core_properties.author = author

    for style, text in paras:
        if style:
            d.add_paragraph(text, style=style)
        else:
            d.add_paragraph(text)

    p = tmp_path / "in.docx"
    d.save(str(p))
    return p


class TestManifest:
    def test_manifest_declares_no_events(self) -> None:
        """Manifest declares empty events tuple."""
        plugin = DocxDownloadSourcePlugin()
        assert plugin.manifest.events == ()


_CLAIMS_TABLE = [
    ("https://example.com/a.docx", True),
    ("https://example.com/a.DOCX", True),
    ("https://example.com/a.Docx?dl=1", True),
    ("https://example.com/a.docx#frag", True),
    ("https://example.com/a.docx?dl=1#frag", True),
    ("https://user@example.com:8080/dir/a.docx", True),
    ("https://example.com/download?file=a.docx", False),
    ("https://example.com/?a.docx", False),
    ("https://example.com/view#a.docx", False),
    ("https://example.docx", False),
    ("https://example.docx?x=1", False),
    ("https://example.docx/", False),
    ("https://example.com/a.docx/", False),
    ("https://example.com/a.docx;v=1", False),
    ("https://example.com/?file=a.docx&x=1", False),
    ("https://example.com/a%2Edocx", False),
    ("https://example.com/a.docx.html", False),
    ("https://example.com/a.docx%20", False),
]


@pytest.mark.parametrize(("url", "claimed"), _CLAIMS_TABLE)
def test_the_manifest_pattern_claims_exactly_what_the_class_claims(url: str, claimed: bool) -> None:
    """The core's out-of-process claim (a regex search of ``url_patterns``) equals the class's own rule."""
    (pattern,) = DocxDownloadSourcePlugin.manifest.url_patterns
    assert DocxDownloadSourcePlugin().claims(url) is claimed
    assert (re.search(pattern, url) is not None) is claimed


class TestClaimsFormat:
    def test_claims_format_docx(self) -> None:
        """claims_format('docx') returns True."""
        plugin = DocxDownloadSourcePlugin()
        assert plugin.claims_format("docx") is True

    def test_claims_format_non_docx(self) -> None:
        """claims_format('epub') returns False."""
        plugin = DocxDownloadSourcePlugin()
        assert plugin.claims_format("epub") is False


class TestPullConvertsAndMaps:
    @responses.activate
    def test_pull_converts_and_maps(self, tmp_path: Path) -> None:
        """pull() downloads DOCX, converts to EPUB, reads metadata, and returns BookPatch."""
        docx_path = _make_docx(
            tmp_path,
            paras=[
                ("Heading 1", "Chapter One"),
                (None, "Body text"),
            ],
            title="My Book",
            author="Jane",
        )
        docx_bytes = docx_path.read_bytes()

        url = "https://x.com/download.docx"
        responses.add(
            responses.GET,
            url,
            body=docx_bytes,
            status=200,
            content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )

        session = requests.Session()
        plugin = DocxDownloadSourcePlugin(session=session)

        work_dir = tmp_path / "work"
        work_dir.mkdir()

        ctx = FakeContext()

        patch = plugin.pull(url, work_dir, None, ctx)

        assert isinstance(patch, BookPatch)
        assert patch.upsert is True
        assert patch.book_id == make_book_id(url)
        assert patch.fields["title"] == "My Book"
        assert patch.fields["author"] == "Jane"
        assert patch.fields["output_filename"] == "Jane/download.epub"
        assert patch.fields["format"] == "docx"

        # Verify the converted EPUB exists and is readable
        epub_path = work_dir / patch.fields["output_filename"]
        assert epub_path.exists()
        epub_doc = EpubDocument.open(epub_path)
        meta = epub_doc.read_metadata()
        assert meta.title == "My Book"

        # Verify the downloaded DOCX is cleaned up
        assert not (work_dir / "download.docx").exists()


class TestPullConversionError:
    @responses.activate
    def test_pull_conversion_error_maps_to_sourcepullerror(self, tmp_path: Path) -> None:
        """pull() raises SourcePullError on DOCX conversion failure."""
        from ebookerr_sdk.spi import SourcePullError

        url = "https://x.com/download.docx"
        responses.add(
            responses.GET,
            url,
            body=b"not a docx",
            status=200,
            content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )

        session = requests.Session()
        plugin = DocxDownloadSourcePlugin(session=session)

        work_dir = tmp_path / "work"
        work_dir.mkdir()

        ctx = FakeContext()

        with pytest.raises(SourcePullError, match="DOCX conversion failed"):
            plugin.pull(url, work_dir, None, ctx)


class TestPullHttpError:
    @responses.activate
    def test_pull_http_error_raises_sourcepullerror(self, tmp_path: Path) -> None:
        """pull() raises SourcePullError on non-200 status."""
        from ebookerr_sdk.spi import SourcePullError

        url = "https://x.com/download.docx"
        responses.add(responses.GET, url, status=404)

        session = requests.Session()
        plugin = DocxDownloadSourcePlugin(session=session)

        work_dir = tmp_path / "work"
        work_dir.mkdir()

        ctx = FakeContext()

        with pytest.raises(SourcePullError, match="HTTP 404"):
            plugin.pull(url, work_dir, None, ctx)

    @responses.activate
    def test_pull_403_explains_the_refusal(self, tmp_path: Path) -> None:
        """pull() explains a 403 refusal instead of a bare status code."""
        from ebookerr_sdk.spi import SourcePullError

        url = "https://x.com/download.docx"
        responses.add(responses.GET, url, status=403)

        session = requests.Session()
        plugin = DocxDownloadSourcePlugin(session=session)

        work_dir = tmp_path / "work"
        work_dir.mkdir()

        ctx = FakeContext()

        with pytest.raises(SourcePullError) as exc_info:
            plugin.pull(url, work_dir, None, ctx)
        assert str(exc_info.value) == (
            "download refused: HTTP 403 — the site denied access"
            " (the sign-in may have expired, or access to this item has ended)"
        )


class TestCheckForUpdateDelegates:
    @responses.activate
    def test_check_for_update_delegates(self) -> None:
        """plugin's check_for_update with prior → delegates to helper."""
        from ebookerr_sdk.spi import CustomValueView

        url = "https://x.com/book.docx"

        # Mock a HEAD response
        responses.add(
            responses.HEAD,
            url,
            status=200,
            headers={"ETag": '"abc"'},
        )

        session = requests.Session()
        plugin = DocxDownloadSourcePlugin(session=session)

        # Create prior with stored etag
        prior = BookView(
            book_id=make_book_id(url),
            title="Test Book",
            author="Test Author",
            story_url=url,
            output_filename="test.docx",
            num_chapters=1,
            status=None,
            rating=None,
            cover_ref=None,
            external=None,
            progress=None,
            custom_values={
                "etag": CustomValueView(
                    value="abc",
                    value_type="str",
                    updated_at=None,
                ),
            },
        )

        result = plugin.check_for_update(url, prior=prior)

        # Verify it returns an UpdateCheck with needs_update=False (etag matches)
        from ebookerr_sdk.spi import UpdateCheck

        assert isinstance(result, UpdateCheck)
        assert result.needs_update is False
        assert result.error is None

        # Verify exactly one HEAD call was made
        assert len(responses.calls) == 1
        assert responses.calls[0].request.method == "HEAD"


class TestTitleAndAuthorMetadataPrecedence:
    """Test title and author fallback logic per EXP-224."""

    @responses.activate
    @pytest.mark.pins("EXP-224")
    @pytest.mark.real_impl("docx_download_source.docx_to_epub.convert_docx_to_epub")
    def test_title_and_author_fall_back_to_the_document_before_the_url_stem(
        self, tmp_path: Path, caplog: Any
    ) -> None:
        """Real DOCX converter: URL stem chapter179 + first_line → uses first_line."""
        from docx_download_source.docx_to_epub import convert_docx_to_epub
        from ebookerr_sdk.download.document import download_convert_stage

        DOCX_CONTENT_TYPE = (
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )

        body = (
            Path(__file__).resolve().parent / "fixtures" / "Three Square Meals - Chapter 179.docx"
        ).read_bytes()
        url = "https://ex.com/honest/chapter179.docx"
        responses.add(
            responses.GET,
            url,
            body=body,
            status=200,
            headers={"Content-Type": DOCX_CONTENT_TYPE},
        )

        session = requests.Session()
        work_dir = tmp_path / "work"
        work_dir.mkdir()

        ctx = FakeContext()

        with caplog.at_level(logging.DEBUG):
            patch = download_convert_stage(
                url,
                work_dir,
                None,
                ctx,
                session=session,
                auth_headers=ctx.auth_headers(url),
                timeout_s=300.0,
                namespace="docx_download_source",
                source_suffix=".docx",
                fmt="docx",
                convert=convert_docx_to_epub,
                log_label="DOCX",
            )

        assert patch.fields["title"] == "Three Square Meals Ch. 179"
        assert patch.fields["author"] == "Unknown"
        assert patch.fields["output_filename"].startswith("Unknown/")
        debug_logs = [r.message for r in caplog.records if r.levelno == logging.DEBUG]
        assert any(
            "Discarded junk document author 'Windows User'" in r.message
            for r in caplog.records
            if r.levelno == logging.DEBUG
        ), f"Expected debug log not found. Debug logs: {debug_logs}"
        assert any(
            "Title taken from the document's first line" in r.message
            for r in caplog.records
            if r.levelno == logging.DEBUG
        ), f"Expected 'Title taken from the document' log not found. Debug logs: {debug_logs}"
